"""Load HR policy and noise documents from the local data/ folder.

The source is the filesystem for now. When a remote store is chosen later,
replace `_iter_source_files` — `load_documents` and the parsers stay the same.
"""

from __future__ import annotations

import re
from pathlib import Path

from langchain_core.documents import Document

from hr_assistant.config import DATA_DIR, HR_POLICY_CATEGORIES

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".pptx"}

_CATEGORY_RE = re.compile(
    r"^(?:Policy\s+)?Category:\s*(.+)$",
    re.IGNORECASE | re.MULTILINE,
)


def load_documents(data_dir: Path | None = None) -> list[Document]:
    """Return one LangChain Document per source file under `data_dir` (recursive)."""
    return _load(data_dir, noise_only=None)


def load_hr_documents(data_dir: Path | None = None) -> list[Document]:
    """HR policies only: files in data/ that are not under data/noise/."""
    return _load(data_dir, noise_only=False)


def load_noise_documents(data_dir: Path | None = None) -> list[Document]:
    """Non-HR corpus under data/noise/ (pdf, docx, pptx, txt)."""
    return _load(data_dir, noise_only=True)


def _load(data_dir: Path | None, noise_only: bool | None) -> list[Document]:
    root = Path(data_dir) if data_dir is not None else DATA_DIR
    if not root.is_dir():
        raise FileNotFoundError(f"Data directory not found: {root}")

    documents: list[Document] = []
    for path in _iter_source_files(root):
        is_noise = "noise" in path.relative_to(root).parts
        if noise_only is True and not is_noise:
            continue
        if noise_only is False and is_noise:
            continue
        text = _extract_text(path)
        if not text.strip():
            continue
        documents.append(
            Document(
                page_content=text,
                metadata=_build_metadata(path, root, text),
            )
        )
    return documents


def _iter_source_files(root: Path) -> list[Path]:
    files = [
        path
        for path in sorted(root.rglob("*"))
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and "processed" not in path.relative_to(root).parts
    ]
    return files


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pptx":
        return _extract_pptx(path)
    return ""


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    blocks: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    blocks.append(text)
    return "\n".join(blocks)


def _build_metadata(path: Path, root: Path, text: str) -> dict:
    category = extract_policy_category(text)
    relative = path.relative_to(root).as_posix()
    return {
        "source": relative,
        "filename": path.name,
        "file_type": path.suffix.lower().lstrip("."),
        "category": category,
        "policy_category": category,
        "is_hr": category in HR_POLICY_CATEGORIES,
    }


def extract_policy_category(text: str) -> str:
    match = _CATEGORY_RE.search(text)
    if not match:
        return ""
    return match.group(1).strip()


if __name__ == "__main__":
    docs = load_documents()
    print(f"Loaded {len(docs)} documents from {DATA_DIR}")
    for doc in docs:
        meta = doc.metadata
        print(
            f"  {meta['source']:40} category={meta['category'] or '-':20} "
            f"hr={meta['is_hr']} chars={len(doc.page_content)}"
        )
